"""Generate deterministic fictional fixtures; no private source documents."""
import json
import sys
from pathlib import Path

root = Path(__file__).parent / 'fixtures'
cases, manifest = [], []
extensions = ['txt', 'md', 'docx', 'pdf'] * 2
facts = []
for i in range(8):
    name = f'星桥X{i+1}'
    fact = f'{name}整机保修期为{12+i*6}个月，核心部件保修期为{24+i*6}个月。'
    facts.append(fact)
for i, ext in enumerate(extensions):
    name = f'星桥X{i+1}'
    manifest.append({'file': f'product-{i+1}.{ext}', 'title': name, 'fact': facts[i]})
    split = 'calibration' if i < 4 else 'acceptance'
    for category, query, history, expected in [
        ('fact', f'{name}整机保修期是几个月？', [], [facts[i]]),
        ('semantic', f'购买{name}后，整台机器坏了厂家免费负责多久？', [], [facts[i]]),
        ('followup', '它的核心部件呢？', [{'role': 'user', 'content': f'我想了解{name}的保修政策'}], [facts[i]]),
        ('multi', f'比较{name}和星桥X{(i+1)%8+1}的整机及核心部件保修期。', [], [facts[i], facts[(i+1)%8]]),
        ('no_answer', f'{name}的内部量子芯片序列号是什么？', [], []),
    ]:
        cases.append({'id': f'{category}-{i+1}', 'category': category, 'split': split, 'query': query,
                      'history': history, 'expectedEvidence': expected, 'evidenceLocation': '第1节 保修条款' if expected else None})

for name, generated in [('manifest.json', manifest), ('cases.json', cases)]:
    expected = json.loads((root.parent / name).read_text(encoding='utf8'))
    if generated != expected:
        raise RuntimeError(f'{name} 与固定评测定义不一致；请先核对题目划分，不能静默改写验收集')

if '--verify-definitions' in sys.argv:
    print(f'Verified {len(manifest)} documents and {len(cases)} cases')
    raise SystemExit(0)

from docx import Document
from docx.shared import RGBColor
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from PIL import Image, ImageDraw, ImageFont

root.mkdir(exist_ok=True)
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
for i, ext in enumerate(extensions):
    name = f'星桥X{i+1}'
    sections = [f'{name}产品政策（虚构测试资料）', '第1节 保修条款\n' + facts[i]]
    for j in range(2, 17):
        sections.append(f'第{j}节 归档记录\n' + ''.join(
            f'{name}第{j}批第{k}项记录只描述包装检查：核对外箱、标签、缓冲材料与运输记录，不改变保修政策。'
            for k in range(1, 8)))
    sections.append(f'第17节 售后表\n项目 | 规定\n退款到账 | {i+2}个工作日\n维修受理 | 工作日上午九点\n这些规定只适用于{name}，不能套用到其他型号。')
    text = '\n\n'.join(sections)
    file = root / f'product-{i+1}.{ext}'
    if ext in ['txt', 'md']:
        file.write_text(text, encoding='utf8')
    elif ext == 'docx':
        doc = Document()
        for style in ['Title', 'Heading 1']:
            doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
        doc.add_paragraph(sections[0], 'Title')
        for section in sections[1:]:
            heading, body = section.split('\n', 1)
            doc.add_paragraph(heading, 'Heading 1')
            doc.add_paragraph(body)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '型号'
        table.cell(0, 1).text = name
        doc.save(file)
    else:
        pdf = canvas.Canvas(str(file))
        y = 800
        pdf.setFont('STSong-Light', 11)
        for line in text.splitlines():
            for start in range(0, max(1, len(line)), 43):
                if y < 50:
                    pdf.showPage(); pdf.setFont('STSong-Light', 11); y = 800
                pdf.drawString(45, y, line[start:start+43]); y -= 17
        pdf.save()
(root / 'empty.txt').write_text('', encoding='utf8')
(root / 'broken.pdf').write_bytes(b'not a PDF')
img = Image.new('RGB', (900, 180), 'white')
ImageDraw.Draw(img).text((20, 40), 'Scanned fictional document: warranty 12 months.', fill='black', font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 25))
pdf = canvas.Canvas(str(root / 'scan.pdf'))
from reportlab.lib.utils import ImageReader
pdf.drawImage(ImageReader(img), 30, 650, width=540, height=108)
pdf.save()
print(f'Generated {len(manifest)} documents and {len(cases)} cases')
